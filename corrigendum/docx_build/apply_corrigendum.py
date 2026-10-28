"""Apply the minimal flat10MIP corrigendum to the converted paper as tracked changes."""
import sys, csv
from docx import Document
from docx.shared import Emu
from docx.oxml.ns import qn
from track import replace, replace_all, append, swap_drawing, ptext

src, out, csv_path, figdir = sys.argv[1:5]
d = Document(src)
NB = ' '; M = '□'
paras = list(d.element.body.iter(qn('w:p')))
def para(snippet):
    hits = [p for p in paras if snippet in ptext(p)]
    assert len(hits) == 1, f'{len(hits)} paragraphs contain: {snippet[:60]}'
    return hits[0]
def edit(old, new): replace(para(old), old, new)

edit('reduced by 30 % compared', 'reduced by about 40 % compared')                                      # 1
edit(f'TNZ is calculated as a 21-year average around year 150 in esm-flat10-cdr (i.e., 50{NB}years after branching from esm-flat10) minus the expected temperature at net zero using cumulative emissions proportionality ({M}, which is 1.25 times the esm-flat10-derived TCRE; see Fig.{NB}2).',
     f'TNZ is calculated as the 20-year mean temperature in esm-flat10-cdr centred on year 150 (years 140–159, i.e., 50{NB}years after branching from esm-flat10) minus the expected temperature at net zero using cumulative emissions proportionality ({M}, the 20-year mean temperature in esm-flat10 centred on year 125, years 115–134, when cumulative emissions in esm-flat10 also reach 1250 PgC; see Fig.{NB}2).')  # 2
edit('the peak value of 20-year smoothed global mean temperatures', 'the peak value of 11-year smoothed global mean temperatures')  # 3
edit('(but most models peak before)', '(four of the ten ESMs peak before net zero and six after)')       # 4
edit(f'ESMs are also consistent with the relationship of {M} 25 % smaller absolute magnitudes in ZEC50 and ZEC90, albeit with larger scatter.',
     f'ESMs show a larger reduction, with {M} 40 % smaller absolute magnitudes in ZEC50 and ZEC90 (best-fit slopes of 0.59 and 0.57; Fig.{NB}8c, d), albeit with larger scatter.')  # 6
edit('lie outside of the SCM distributions such that we see peak warming significantly before net zero with greater warming', 'lie outside of the SCM distributions with greater warming')  # 6
edit('measured as the warming in esm-flat10-zec between years 100 and years 150/190/200, respectively.',
     'measured as the difference between the 20-year mean warming in esm-flat10-zec centred on years 150/190/200 and the 20-year mean warming in esm-flat10 centred on year 100, respectively.')  # Table 2 caption

rows = {r['model']: r for r in csv.DictReader(open(csv_path))}
name = {'UKESM1.2': 'UKESM'}
cols = {3: 'T100yr', 6: 'ZEC50', 7: 'ZEC90', 8: 'ZEC100', 9: 'tPW', 12: 'TR0'}
approx_tr0 = {'CESM2', 'GISS', 'UKESM'}
def fmt(v, k):
    s = f'{float(v):.0f}' if k == 'tPW' else f'{float(v):.2f}'
    return s.replace('-', '−')
for row in d.tables[1].rows:
    key = name.get(row.cells[0].text.strip(), row.cells[0].text.strip())
    if key not in rows: continue
    for ci, k in cols.items():
        p = row.cells[ci].paragraphs[0]._p
        new = fmt(rows[key][k], k); sup = 'd' if (k == 'TR0' and key in approx_tr0) else None
        old = ptext(p).strip().replace('-', '−')
        if k == 'tPW' and old.endswith('.0'): old = old[:-2]
        if old != new or sup:
            replace_all(p, new, sup)
append(para('Chris Jones (personal communication, 2024)'),
       ' d TR0 approximated from years 291–300 because the esm-flat10-cdr simulation ends in year 300: the mean temperature over those years minus the warming expected from the 45 PgC of cumulative emissions remaining on average (0.045 × T100yr).')

runs = []
for b in d.element.body.iter(qn('a:blip')):
    r = b
    while r.tag != qn('w:r'): r = r.getparent()
    runs.append(r)
assert len(runs) == 13
for n in (7, 8, 13):
    old = runs[n - 1]
    width = Emu(int(old.find('.//' + qn('wp:extent')).get('cx')))
    tmp = d.add_paragraph(); tmp.add_run().add_picture(f'{figdir}/fig{n:02d}.png', width=width)
    swap_drawing(old, tmp._p.find(qn('w:r'))); tmp._p.getparent().remove(tmp._p)
d.save(out); print('applied')
