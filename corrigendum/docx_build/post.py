"""Post-process the pandoc docx: A4 pages, compact full-width tables, Table 2 on a landscape page."""
import sys, copy
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out = sys.argv[1], sys.argv[2]
d = Document(src)
for s in d.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2.2)

def set_table(tbl, size):
    tblPr = tbl._tbl.tblPr
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW'); tblPr.append(w)
    w.set(qn('w:w'), '5000'); w.set(qn('w:type'), 'pct')
    lay = tblPr.find(qn('w:tblLayout'))
    if lay is None:
        lay = OxmlElement('w:tblLayout'); tblPr.append(lay)
    lay.set(qn('w:type'), 'autofit')
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'bottom', 'insideH'):
        b = OxmlElement(f'w:{side}'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), '999999')
        borders.append(b)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    tblPr.append(borders)
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                for r in p.runs: r.font.size = Pt(size)
                for m in p._p.iter(qn('m:r')):
                    pass

body = d.element.body
tables = d.tables
for i, t in enumerate(tables):
    set_table(t, 7 if i == 1 else 8)

# Table 2 on its own landscape section
def sect_pr(landscape):
    base = copy.deepcopy(body.find(qn('w:sectPr')))
    pg = base.find(qn('w:pgSz'))
    if landscape:
        pg.set(qn('w:w'), str(int(Cm(29.7).twips))); pg.set(qn('w:h'), str(int(Cm(21.0).twips))); pg.set(qn('w:orient'), 'landscape')
    else:
        pg.set(qn('w:w'), str(int(Cm(21.0).twips))); pg.set(qn('w:h'), str(int(Cm(29.7).twips)))
        if pg.get(qn('w:orient')): del pg.attrib[qn('w:orient')]
    return base

if len(tables) >= 2:
    t2 = tables[1]._tbl
    cap = t2.getprevious()          # the caption paragraph precedes the table
    before = cap.getprevious()
    pPr = before.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); before.insert(0, pPr)
    pPr.append(sect_pr(False))
    after = t2.getnext()            # table footnote paragraph
    while after is not None and after.tag != qn('w:p'):
        after = after.getnext()
    pPr = after.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); after.insert(0, pPr)
    pPr.append(sect_pr(True))
d.save(out)
